"""Generic Excel parser used by all three Step 6 import kinds.

Each ingest module (glossary / business_rules / column_rules) wraps this
parser with a kind-specific spec: the expected column headers, which are
required, and a per-row validation hook.

Design notes:
- One sheet per upload (the first one). Multi-sheet workbooks are
  flagged in the import error report but only the first sheet is read.
- Headers are matched **case-insensitively** and trimmed of whitespace
  so users typing "Term " or "term" or "TERM" all work.
- Per-row errors are collected, never raised — the import lands in the
  ``error`` status with the full error list in the row, but the user
  can fix the spreadsheet and re-upload.
- Empty rows (every cell is None or empty-string) are skipped silently.
- Synonym-like columns use semicolons as separators (BRD Tables 11/12).
  We accept commas too because users habitually use them.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any, Callable

from openpyxl import load_workbook

logger = logging.getLogger(__name__)


@dataclass
class ColumnSpec:
    """One column expected in the upload."""
    name: str                            # canonical header (lower-case key)
    required: bool = False
    aliases: tuple[str, ...] = ()        # extra header strings that map here
    coerce: Callable[[Any], Any] | None = None  # optional value transform


@dataclass
class ParsedRow:
    row_number: int                      # 1-based, matches the user's view in Excel
    values: dict[str, Any]               # keyed by canonical column name
    errors: list[str] = field(default_factory=list)


@dataclass
class ParseResult:
    rows: list[ParsedRow] = field(default_factory=list)
    file_errors: list[str] = field(default_factory=list)   # whole-file issues
    column_order: list[str] = field(default_factory=list)

    @property
    def error_count(self) -> int:
        return len(self.file_errors) + sum(1 for r in self.rows if r.errors)

    @property
    def ok_rows(self) -> list[ParsedRow]:
        return [r for r in self.rows if not r.errors]

    def to_error_jsonb(self) -> list[dict]:
        """Shape consumed by t_dq_imports.errors (JSONB)."""
        out: list[dict] = []
        for e in self.file_errors:
            out.append({"row": None, "error": e})
        for r in self.rows:
            for e in r.errors:
                out.append({"row": r.row_number, "error": e})
        return out


def split_list_cell(raw: Any) -> list[str]:
    """Split a synonym-list cell (semicolon-preferred, comma-tolerant).
    Returns deduped, stripped, lowercase entries; empty list for empty input."""
    if raw is None:
        return []
    s = str(raw).strip()
    if not s:
        return []
    parts = [p.strip().lower() for p in s.replace(",", ";").split(";")]
    seen: set[str] = set()
    out: list[str] = []
    for p in parts:
        if p and p not in seen:
            seen.add(p)
            out.append(p)
    return out


def parse_excel(
    file_bytes: bytes, columns: list[ColumnSpec], *,
    row_validators: list[Callable[[ParsedRow], None]] | None = None,
    max_rows: int = 10_000,
) -> ParseResult:
    """Parse the first sheet of an .xlsx file against the supplied column
    spec. Returns a ParseResult with one entry per non-empty row plus
    file-level errors. Never raises on bad input."""
    result = ParseResult()

    try:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        result.file_errors.append(f"Could not open as .xlsx: {e}")
        return result

    if not wb.sheetnames:
        result.file_errors.append("Workbook has no sheets")
        return result
    if len(wb.sheetnames) > 1:
        result.file_errors.append(
            f"Workbook has {len(wb.sheetnames)} sheets; only the first "
            f"({wb.sheetnames[0]!r}) is read."
        )

    ws = wb[wb.sheetnames[0]]

    # --- Header row -----------------------------------------------------
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        result.file_errors.append("Header row missing or empty")
        return result

    # Index canonical → 0-based column index, using aliases tolerantly.
    alias_map: dict[str, str] = {}
    for spec in columns:
        for key in (spec.name, *spec.aliases):
            alias_map[key.lower().strip()] = spec.name

    header_to_col: dict[str, int] = {}
    for i, h in enumerate(header_row):
        if h is None:
            continue
        key = str(h).lower().strip()
        canonical = alias_map.get(key)
        if canonical:
            header_to_col[canonical] = i
    result.column_order = list(header_to_col.keys())

    # Required-column check
    missing_required = [
        s.name for s in columns
        if s.required and s.name not in header_to_col
    ]
    if missing_required:
        result.file_errors.append(
            f"Missing required column(s): {', '.join(missing_required)}"
        )
        return result

    # --- Data rows ------------------------------------------------------
    seen_rows = 0
    for r_idx, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2,
    ):
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue
        if seen_rows >= max_rows:
            result.file_errors.append(
                f"Truncated after {max_rows} rows — split the file."
            )
            break

        parsed = ParsedRow(row_number=r_idx, values={})
        for spec in columns:
            idx = header_to_col.get(spec.name)
            raw = row[idx] if idx is not None and idx < len(row) else None
            if isinstance(raw, str):
                raw = raw.strip() or None
            if spec.required and (raw is None or raw == ""):
                parsed.errors.append(f"Missing required value: {spec.name}")
            if spec.coerce is not None and raw is not None:
                try:
                    raw = spec.coerce(raw)
                except Exception as e:  # noqa: BLE001
                    parsed.errors.append(f"{spec.name}: {e}")
                    raw = None
            parsed.values[spec.name] = raw

        if row_validators:
            for v in row_validators:
                try:
                    v(parsed)
                except Exception as e:  # noqa: BLE001
                    parsed.errors.append(f"validator error: {e}")
        result.rows.append(parsed)
        seen_rows += 1

    return result


def write_glossary_template(out: BytesIO) -> BytesIO:
    """Build a downloadable .xlsx template for kind=glossary with the
    Table-11 headers + one example row. Used by GET /imports/template."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Glossary"
    ws.append(["term", "definition", "synonyms", "language"])
    ws.append([
        "customer_email",
        "The primary email address for a customer record.",
        "email; user_email; contact_email; البريد_الإلكتروني",
        "mixed",
    ])
    ws.append([
        "national_id",
        "Saudi Arabian National ID number (10 digits).",
        "saudi_id; id_number; الهوية_الوطنية",
        "mixed",
    ])
    wb.save(out)
    out.seek(0)
    return out
